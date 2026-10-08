"""Core NLP / scoring engine for the AI Resume Analyzer.

Techniques used:
  * Regex based skill extraction with alias matching
  * TF-IDF + cosine similarity between resume and job description (scikit-learn)
  * Rule based ATS checks (sections, contact info, action verbs, quantified impact, length)
  * Weighted score aggregation using NumPy
"""
import io
import re

import numpy as np
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from skills_db import ACTION_VERBS, LEARNING_HINTS, ROLES, SKILLS

# ---------------------------------------------------------------- text input

def extract_text(file_name: str, data: bytes) -> str:
    """Return plain text from a PDF, DOCX or TXT resume."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        import docx
        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)
    return data.decode("utf-8", errors="ignore")


# ----------------------------------------------------------- skill extraction

def _build_patterns():
    patterns = {}
    for category, skills in SKILLS.items():
        for canonical, aliases in skills.items():
            alts = "|".join(re.escape(a) for a in sorted(set(aliases), key=len, reverse=True))
            regex = re.compile(r"(?<![a-z0-9+#])(?:" + alts + r")(?![a-z0-9+#])")
            patterns[canonical] = (category, regex)
    return patterns


_PATTERNS = _build_patterns()


def find_skills(text: str) -> dict:
    """Return {skill: category} for every known skill found in the text."""
    lower = text.lower()
    return {skill: cat for skill, (cat, rx) in _PATTERNS.items() if rx.search(lower)}


# ------------------------------------------------------------- ATS style checks

SECTION_KEYWORDS = {
    "education": ["education", "academic"],
    "skills": ["skills", "technical skills", "technologies"],
    "projects": ["projects", "project work", "academic projects"],
    "experience": ["experience", "internship", "internships", "work history", "employment"],
    "summary": ["summary", "objective", "career objective", "profile", "about me"],
    "certifications": ["certification", "certifications", "certificates", "courses"],
    "achievements": ["achievements", "awards", "accomplishments", "extracurricular",
                     "positions of responsibility"],
}
SECTION_WEIGHTS = {"education": 20, "skills": 20, "projects": 20, "experience": 15,
                   "summary": 10, "certifications": 8, "achievements": 7}


def detect_sections(text: str) -> dict:
    found = {s: False for s in SECTION_KEYWORDS}
    for line in text.splitlines():
        clean = line.strip(" :-•*#\t").lower()
        if not clean or len(clean) > 45:
            continue
        for section, kws in SECTION_KEYWORDS.items():
            if any(clean == kw or clean.startswith(kw + " ") or clean.startswith(kw + ":") for kw in kws):
                found[section] = True
    return found


def detect_contact(text: str) -> dict:
    lower = text.lower()
    return {
        "email": bool(re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)),
        "phone": bool(re.search(r"(\+?\d[\d\s-]{8,13}\d)", text)),
        "linkedin": "linkedin.com" in lower or "linkedin" in lower,
        "github": "github.com" in lower or "github" in lower,
    }


def count_action_verbs(text: str) -> int:
    lower = text.lower()
    return sum(len(re.findall(r"\b" + v + r"\b", lower)) for v in set(ACTION_VERBS))


def count_quantified(text: str) -> int:
    """Count numbers/percentages that suggest measurable impact (ignores years & phone numbers)."""
    cleaned = re.sub(r"\b(?:19|20)\d{2}\b", " ", text)
    cleaned = re.sub(r"\d[\d\s-]{8,}\d", " ", cleaned)
    return len(re.findall(r"\d+(?:\.\d+)?\s?%|\b\d+\+?\b", cleaned))


def length_score(words: int) -> float:
    if 250 <= words <= 800:
        return 100.0
    if words < 250:
        return max(0.0, words / 250 * 100)
    return max(40.0, 100 - (words - 800) / 20)


# ------------------------------------------------------------------- scoring

def tfidf_similarity(resume: str, reference: str) -> float:
    """Cosine similarity between resume and reference text, scaled to 0-100."""
    try:
        vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2), sublinear_tf=True)
        matrix = vec.fit_transform([resume, reference])
        sim = float(cosine_similarity(matrix[0], matrix[1])[0][0])
    except ValueError:
        return 0.0
    # raw cosine for resume-vs-JD is naturally low, so rescale for readability
    return float(np.clip(sim / 0.35, 0, 1) * 100)


WEIGHTS = {"Skill Match": 40, "JD / Role Similarity": 20, "Sections": 15,
           "Contact & Links": 5, "Impact (verbs + numbers)": 10, "Length": 10}


def analyze(resume_text: str, role: str, jd_text: str = "") -> dict:
    role_info = ROLES[role]
    jd_text = (jd_text or "").strip()

    resume_skills = find_skills(resume_text)

    # required skills: from JD if one was pasted (and it has enough known skills), else from role
    jd_skills = find_skills(jd_text) if jd_text else {}
    if len(jd_skills) >= 3:
        required = sorted(jd_skills)
        reference = jd_text
        source = "job description"
    else:
        required = list(role_info["skills"])
        reference = role_info["description"]
        source = f"role: {role}"

    matched = [s for s in required if s in resume_skills]
    missing = [s for s in required if s not in resume_skills]
    extra = sorted(s for s in resume_skills if s not in required)

    skill_score = len(matched) / len(required) * 100 if required else 0.0
    sim_score = tfidf_similarity(resume_text, reference)

    sections = detect_sections(resume_text)
    section_score = sum(SECTION_WEIGHTS[s] for s, ok in sections.items() if ok)

    contact = detect_contact(resume_text)
    contact_score = sum(contact.values()) / len(contact) * 100

    verbs = count_action_verbs(resume_text)
    quant = count_quantified(resume_text)
    impact_score = float(min(100, verbs * 8 + quant * 10))

    words = len(re.findall(r"\w+", resume_text))
    len_score = length_score(words)

    parts = {
        "Skill Match": skill_score,
        "JD / Role Similarity": sim_score,
        "Sections": float(section_score),
        "Contact & Links": contact_score,
        "Impact (verbs + numbers)": impact_score,
        "Length": len_score,
    }
    values = np.array([parts[k] for k in WEIGHTS])
    weights = np.array([WEIGHTS[k] for k in WEIGHTS], dtype=float)
    overall = float(np.average(values, weights=weights))

    result = {
        "overall": round(overall, 1),
        "parts": {k: round(v, 1) for k, v in parts.items()},
        "required": required,
        "matched": matched,
        "missing": missing,
        "extra": extra,
        "resume_skills": resume_skills,
        "sections": sections,
        "contact": contact,
        "verbs": verbs,
        "quantified": quant,
        "words": words,
        "source": source,
    }
    result["suggestions"] = build_suggestions(result)
    result["grade"] = grade(overall)
    return result


def grade(score: float) -> str:
    if score >= 80:
        return "Excellent - placement ready"
    if score >= 65:
        return "Good - minor improvements needed"
    if score >= 45:
        return "Average - work on the gaps below"
    return "Needs significant improvement"


def build_suggestions(r: dict) -> list:
    tips = []
    if r["missing"]:
        top = ", ".join(r["missing"][:5])
        tips.append(f"Add or learn these in-demand skills for this target: {top}. "
                    f"Only list skills you can honestly explain in an interview.")
    for section, ok in r["sections"].items():
        if not ok and section in ("education", "skills", "projects"):
            tips.append(f"Your resume is missing a clear '{section.title()}' section (ATS systems look for it).")
    if not r["sections"]["summary"]:
        tips.append("Add a 2-3 line Career Objective / Summary tailored to the target role.")
    if not r["sections"]["certifications"]:
        tips.append("Add relevant certifications / online courses (Coursera, NPTEL, Udemy, etc.).")
    if not r["contact"]["linkedin"]:
        tips.append("Add your LinkedIn profile link.")
    if not r["contact"]["github"]:
        tips.append("Add your GitHub link with project READMEs - recruiters check it.")
    if not r["contact"]["email"] or not r["contact"]["phone"]:
        tips.append("Make sure a professional email and phone number are clearly visible.")
    if r["verbs"] < 5:
        tips.append("Start bullet points with strong action verbs (Developed, Built, Analyzed, Optimized...).")
    if r["quantified"] < 3:
        tips.append("Quantify achievements - e.g. 'improved accuracy to 92%', 'built dashboard used by 50+ users'.")
    if r["words"] < 250:
        tips.append("Resume looks too short; add more detail about projects and skills (aim for 1 page, ~300-600 words).")
    elif r["words"] > 800:
        tips.append("Resume is long; keep it to one page for a fresher.")
    if r["parts"]["JD / Role Similarity"] < 40:
        tips.append("Use more keywords from the job description / role in your project and skill descriptions.")
    if not tips:
        tips.append("Great resume! Tailor the keywords for each company and keep proofreading.")
    return tips


def learning_plan(missing: list) -> list:
    out = []
    for skill in missing:
        hint = LEARNING_HINTS.get(
            skill,
            f"Search '{skill} for beginners' (YouTube / freeCodeCamp / official docs) and build one mini project.")
        out.append((skill, hint))
    return out


def text_report(role: str, r: dict) -> str:
    lines = [
        "AI RESUME ANALYZER - REPORT",
        "=" * 40,
        f"Target: {role}   (skills compared against {r['source']})",
        f"Overall score: {r['overall']}/100  -> {r['grade']}",
        "",
        "Score breakdown:",
    ]
    lines += [f"  - {k}: {v}" for k, v in r["parts"].items()]
    lines += ["", "Matched skills: " + (", ".join(r["matched"]) or "none"),
              "Missing skills: " + (", ".join(r["missing"]) or "none"),
              "Other skills found: " + (", ".join(r["extra"]) or "none"),
              "", "Suggestions:"]
    lines += [f"  {i}. {t}" for i, t in enumerate(r["suggestions"], 1)]
    lines += ["", "Learning plan:"]
    lines += [f"  - {s}: {h}" for s, h in learning_plan(r["missing"])]
    return "\n".join(lines)
